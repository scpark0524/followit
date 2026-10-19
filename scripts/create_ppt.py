#!/usr/bin/env python3
"""Action Keeper (가칭) 중간발표 PPT 생성 스크립트"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── 설정 ──
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
BG_DARK = RGBColor(0x15, 0x15, 0x15)
BG_SECTION = RGBColor(0xBE, 0xB4, 0xE8)  # 보라색 (교수님 스타일)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
PINK = RGBColor(0xFF, 0x33, 0x77)
LIGHT_GRAY = RGBColor(0xAA, 0xAA, 0xAA)
CARD_GREEN = RGBColor(0xE8, 0xF0, 0xD8)
CARD_PURPLE = RGBColor(0xD8, 0xD0, 0xF0)
CARD_YELLOW = RGBColor(0xF0, 0xEC, 0xD8)
CARD_MINT = RGBColor(0xD8, 0xF0, 0xE8)
DARK_CARD = RGBColor(0x22, 0x22, 0x22)

prs = Presentation()
prs.slide_width = SLIDE_WIDTH
prs.slide_height = SLIDE_HEIGHT


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, font_size=18,
             color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Arial"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multiline(slide, left, top, width, height, lines, font_size=16,
                  color=WHITE, line_spacing=1.5, bold=False):
    """lines: list of (text, optional_overrides_dict)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            text, overrides = item, {}
        else:
            text, overrides = item[0], item[1] if len(item) > 1 else {}

        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(overrides.get("size", font_size))
        p.font.color.rgb = overrides.get("color", color)
        p.font.bold = overrides.get("bold", bold)
        p.font.name = overrides.get("font", "Arial")
        p.space_after = Pt(font_size * (line_spacing - 1))
        p.alignment = overrides.get("align", PP_ALIGN.LEFT)
    return txBox


def add_card(slide, left, top, width, height, bg_color,
             title="", body="", title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    shape.shadow.inherit = False

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(16)
    tf.margin_top = Pt(14)
    tf.margin_right = Pt(16)
    tf.margin_bottom = Pt(14)

    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(title_size)
        p.font.bold = True
        p.font.color.rgb = BLACK
        p.font.name = "Arial"
    if body:
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(body_size)
        p2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        p2.font.name = "Arial"
        p2.space_before = Pt(6)
    return shape


def add_table(slide, left, top, width, rows_data, col_widths=None,
              header_bg=BLACK, header_fg=WHITE, body_bg=BG_DARK, body_fg=WHITE):
    """rows_data: list of lists. First row = header."""
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.5 * n_rows))
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w

    for r, row in enumerate(rows_data):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = cell_text
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.name = "Arial"
                if r == 0:
                    p.font.bold = True
                    p.font.color.rgb = header_fg
                else:
                    p.font.color.rgb = body_fg
            # Cell fill
            cell_fill = cell.fill
            cell_fill.solid()
            if r == 0:
                cell_fill.fore_color.rgb = header_bg
            else:
                cell_fill.fore_color.rgb = body_bg
    return table_shape


# ════════════════════════════════════════════
# S1. 표지
# ════════════════════════════════════════════
s1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_bg(s1, BG_DARK)

add_text(s1, Inches(0.8), Inches(1.5), Inches(8), Inches(0.5),
         "KAIST IMMS · 2026 가을학기 · 인공지능 비즈니스 전략",
         font_size=15, color=LIGHT_GRAY)

add_text(s1, Inches(0.8), Inches(2.3), Inches(11), Inches(1.5),
         "회의 이후를 책임지는 AI",
         font_size=52, color=WHITE, bold=True)

add_text(s1, Inches(0.8), Inches(3.8), Inches(11), Inches(0.8),
         "Action Keeper (가칭)",
         font_size=28, color=PINK, bold=False)

add_text(s1, Inches(0.8), Inches(4.8), Inches(11), Inches(0.8),
         "녹음된 회의를 LLM이 분석해 할 일 → 담당자 → 일정을 만들고,\n일이 끝날 때까지 쫓아다니는 n8n 워크플로",
         font_size=18, color=LIGHT_GRAY)

add_text(s1, Inches(0.8), Inches(6.2), Inches(11), Inches(0.5),
         "팀원: OOO · OOO · OOO · OOO",
         font_size=14, color=LIGHT_GRAY)


# ════════════════════════════════════════════
# S2. 문제 ① - 핵심 페인
# ════════════════════════════════════════════
s2 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s2, BG_DARK)

add_text(s2, Inches(0.8), Inches(0.5), Inches(3), Inches(0.4),
         "문제 ①", font_size=14, color=PINK)

add_text(s2, Inches(0.8), Inches(0.9), Inches(11), Inches(1.0),
         "회의록 AI는 넘쳐나는데,\n왜 회의는 여전히 회의에서 끝나나?",
         font_size=36, color=WHITE, bold=True)

# 기준값 훅
add_card(s2, Inches(0.8), Inches(2.5), Inches(5), Inches(1.4), DARK_CARD,
         "", "")
add_text(s2, Inches(1.2), Inches(2.6), Inches(4), Inches(0.6),
         "○○%", font_size=48, color=PINK, bold=True)
add_text(s2, Inches(1.2), Inches(3.3), Inches(4.2), Inches(0.5),
         "지난 한 달 팀 회의에서 정한 일 중 기한 안에 끝난 비율",
         font_size=14, color=LIGHT_GRAY)

add_text(s2, Inches(6.5), Inches(2.6), Inches(0.3), Inches(0.5),
         "→", font_size=28, color=LIGHT_GRAY)

add_card(s2, Inches(7.2), Inches(2.5), Inches(5.3), Inches(1.4), DARK_CARD,
         "", "")
add_text(s2, Inches(7.6), Inches(2.65), Inches(4.5), Inches(0.5),
         "결정은 됐는데, 실행이 안 된다", font_size=20, color=WHITE, bold=True)
add_text(s2, Inches(7.6), Inches(3.2), Inches(4.8), Inches(0.5),
         "PM이 회의록 정리 → 태스크 등록 → 진행 확인을 수작업으로 하고 있음",
         font_size=13, color=LIGHT_GRAY)

# 근본 원인 3가지
add_text(s2, Inches(0.8), Inches(4.3), Inches(5), Inches(0.4),
         "근본 원인 3가지", font_size=16, color=WHITE, bold=True)

add_card(s2, Inches(0.8), Inches(4.9), Inches(3.6), Inches(1.8), CARD_GREEN,
         "① 주인 불명확", '"그거 누가 하기로 했죠?"\n결정은 했는데 담당자를 안 정함')
add_card(s2, Inches(4.7), Inches(4.9), Inches(3.6), Inches(1.8), CARD_PURPLE,
         "② 기한 모호", '"다음 주쯤", "가능한 빨리"\n구체적 날짜 없이 끝남')
add_card(s2, Inches(8.6), Inches(4.9), Inches(3.6), Inches(1.8), CARD_YELLOW,
         "③ 팔로업 부재", '다음 회의에서 아무도 확인 안 하고\n새 안건으로 넘어감')


# ════════════════════════════════════════════
# S3. 문제 ② - 기존 서비스의 한계
# ════════════════════════════════════════════
s3 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s3, BG_DARK)

add_text(s3, Inches(0.8), Inches(0.5), Inches(3), Inches(0.4),
         "문제 ②", font_size=14, color=PINK)

add_text(s3, Inches(0.8), Inches(0.9), Inches(11), Inches(1.0),
         "기존 AI 회의 도구는 '등록'에서 멈춘다",
         font_size=36, color=WHITE, bold=True)

# 비교 표
table_data = [
    ["단계", "Circleback · Otter 등", "Action Keeper (가칭)"],
    ["녹음 → 전사 → 요약", "✅", "✅  (기존 템플릿 활용)"],
    ["액션아이템 추출 + 등록", "✅", "✅"],
    ["R&R 추천 (역할·업무량 기반)", "❌", "✅"],
    ["일정 역산 + 선후관계", "❌", "✅"],
    ["팔로업 대화 + 답변 해석", "❌", "✅"],
    ["지연 감지 → 파급 관리", "❌", "✅"],
    ["다음 회의 안건 자동 생성", "❌", "✅"],
]

tbl = add_table(s3, Inches(0.8), Inches(2.2), Inches(11.5), table_data,
                col_widths=[Inches(3.8), Inches(3.5), Inches(4.2)])

# 핵심 메시지
add_text(s3, Inches(0.8), Inches(6.2), Inches(11), Inches(0.6),
         "기존 도구는 회의를 '기록'합니다.  우리는 회의를 '완결'합니다.",
         font_size=20, color=PINK, bold=True)

add_text(s3, Inches(0.8), Inches(6.7), Inches(11), Inches(0.4),
         "추출·등록은 기존 템플릿을 활용하고, 우리는 확인 · 추적 · 완결에 집중합니다.",
         font_size=14, color=LIGHT_GRAY)


# ════════════════════════════════════════════
# S4. 솔루션 - 5단계 워크플로
# ════════════════════════════════════════════
s4 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s4, BG_DARK)

add_text(s4, Inches(0.8), Inches(0.5), Inches(3), Inches(0.4),
         "솔루션", font_size=14, color=PINK)

add_text(s4, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7),
         "녹음부터 다음 회의까지, 끝나지 않는 고리",
         font_size=36, color=WHITE, bold=True)

# 5단계 카드 (가로 배치)
stages = [
    ("STAGE 1", "기록", "녹음 → STT\n화자 구분", CARD_GREEN),
    ("STAGE 2", "추출", "LLM → 액션·담당자\n기한·근거·확신도", CARD_GREEN),
    ("STAGE 3", "배정 + 확인", "R&R 추천 + 타임라인\nHITL 승인 → 등록", CARD_PURPLE),
    ("STAGE 4", "추적", "D-2 팔로업 DM\n답변 해석 → 분기", CARD_PURPLE),
    ("STAGE 5", "순환", "미완료 → 안건\n다음 회의로 연결", CARD_PURPLE),
]

card_w = Inches(2.2)
gap = Inches(0.15)
start_x = Inches(0.8)

for i, (label, title, body, color) in enumerate(stages):
    x = start_x + (card_w + gap) * i
    y = Inches(2.0)
    add_card(s4, x, y, card_w, Inches(2.2), color, f"{label}\n{title}", body,
             title_size=14, body_size=12)
    # 화살표 (마지막 제외)
    if i < 4:
        add_text(s4, x + card_w, y + Inches(0.8), gap + Inches(0.1), Inches(0.4),
                 "→", font_size=20, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# 순환 화살표 표시
add_text(s4, Inches(5), Inches(4.4), Inches(4), Inches(0.5),
         "↻  STAGE 5 → 다음 회의 녹음 → STAGE 1  (닫힌 루프)",
         font_size=16, color=PINK, bold=True, alignment=PP_ALIGN.CENTER)

# AI가 꼭 필요한 이유
add_text(s4, Inches(0.8), Inches(5.2), Inches(5), Inches(0.4),
         "AI가 꼭 필요한 이유 (규칙으로는 안 되는 3가지)",
         font_size=15, color=WHITE, bold=True)

reasons = [
    '"그건 제가 한번 볼게요" → 액션아이템인지 판단 (키워드 규칙 불가)',
    '"중간발표 전까지는 돼야죠" → 실제 날짜로 변환 (맥락 이해 필요)',
    '"자료가 아직 안 와서요" → 막힘(blocker) 상태 분류 (자유 답변 해석)',
]
add_multiline(s4, Inches(0.8), Inches(5.7), Inches(11), Inches(1.5),
              [(f"  {i+1}.  {r}", {"size": 14}) for i, r in enumerate(reasons)],
              color=LIGHT_GRAY, line_spacing=1.6)


# ════════════════════════════════════════════
# S5. 성공 기준
# ════════════════════════════════════════════
s5 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s5, BG_DARK)

add_text(s5, Inches(0.8), Inches(0.5), Inches(3), Inches(0.4),
         "성공 기준", font_size=14, color=PINK)

add_text(s5, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7),
         "3층으로 측정한다: 품질 · 성과 · 효율",
         font_size=36, color=WHITE, bold=True)

# 3층 카드
add_card(s5, Inches(0.8), Inches(2.2), Inches(3.6), Inches(2.8), CARD_GREEN,
         "품질 (AI가 잘 뽑는가)", "")
add_multiline(s5, Inches(1.2), Inches(2.9), Inches(3), Inches(2),
              [("재현율 ≥ 90%", {"size": 14, "color": BLACK}),
               ("정밀도 ≥ 85%", {"size": 14, "color": BLACK}),
               ("담당자/기한 정확도 ≥ 85%", {"size": 14, "color": BLACK}),
               ("승인 시 수정률 ≤ 20%", {"size": 14, "color": BLACK})],
              line_spacing=1.8)

add_card(s5, Inches(4.7), Inches(2.2), Inches(3.6), Inches(2.8), CARD_PURPLE,
         "성과 (일이 실제로 끝나는가)", "")
add_multiline(s5, Inches(5.1), Inches(2.9), Inches(3), Inches(2),
              [("기한 내 완료율 +30%p ⭐", {"size": 14, "color": BLACK, "bold": True}),
               ("주인 없는 액션 = 0건", {"size": 14, "color": BLACK}),
               ("지연 사전 감지율 ≥ 70%", {"size": 14, "color": BLACK})],
              line_spacing=1.8)

add_card(s5, Inches(8.6), Inches(2.2), Inches(3.6), Inches(2.8), CARD_YELLOW,
         "효율 (사람이 편해지는가)", "")
add_multiline(s5, Inches(9.0), Inches(2.9), Inches(3), Inches(2),
              [("회의 후 정리 시간", {"size": 14, "color": BLACK}),
               ("30분 → 5분 (승인만)", {"size": 14, "color": BLACK, "bold": True})],
              line_spacing=1.8)

# 실험 설계
add_text(s5, Inches(0.8), Inches(5.5), Inches(11), Inches(0.4),
         "실험 설계", font_size=16, color=WHITE, bold=True)

add_multiline(s5, Inches(0.8), Inches(6.0), Inches(11), Inches(1.2),
              [("  •  전후 비교: 도입 전 2주(수작업) vs 도입 후 6주", {"size": 14}),
               ("  •  교란 변수 인정: 학기 후반 마감 압박 → 시기 무관 지표 병행 (AI만 잡아낸 액션 수)", {"size": 14}),
               ("  •  이 팀의 회의 자체가 실험장 → 학기 내내 데이터 축적", {"size": 14})],
              color=LIGHT_GRAY, line_spacing=1.5)


# ════════════════════════════════════════════
# S6. 수행 계획
# ════════════════════════════════════════════
s6 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s6, BG_DARK)

add_text(s6, Inches(0.8), Inches(0.5), Inches(3), Inches(0.4),
         "수행 계획", font_size=14, color=PINK)

add_text(s6, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7),
         "이 팀 자체가 실험장이다",
         font_size=36, color=WHITE, bold=True)

# 타임라인 표
plan_data = [
    ["기간", "마일스톤", "산출물"],
    ["10/24", "중간 발표 (오늘)", "문제 정의 + 설계"],
    ["10/25 ~ 11/7", "STAGE 1~2 구현", "녹음→추출→승인→등록 (최소 기능)"],
    ["11/8 ~ 11/21", "STAGE 3 + 실데이터 수집 시작", "R&R 추천 + 타임라인 + 팀 회의 적용"],
    ["11/22 ~ 12/5", "STAGE 4~5 순환 완성", "팔로업 + 안건 + 대시보드"],
    ["12/6 ~ 12/18", "데이터 축적 + 결과 분석", "이행률 데이터 + 성과 리포트"],
    ["12/19", "기말 발표", "라이브 데모 + 실측 결과"],
]

add_table(s6, Inches(0.8), Inches(2.2), Inches(11.5), plan_data,
          col_widths=[Inches(2.2), Inches(4.5), Inches(4.8)])

# 핵심 포인트
add_text(s6, Inches(0.8), Inches(5.8), Inches(11), Inches(0.5),
         "핵심: 11월 초부터 팀 회의에 실적용 → 5~6주치 실데이터 확보",
         font_size=18, color=PINK, bold=True)

# 팀 역할
add_text(s6, Inches(0.8), Inches(6.4), Inches(11), Inches(0.4),
         "팀 역할 분담", font_size=15, color=WHITE, bold=True)

role_data = [
    ["역할", "담당 업무"],
    ["PM / 발표", "일정 관리, 실험 운영, 발표 스토리"],
    ["n8n 워크플로", "5개 서브 워크플로 구축, API 연동"],
    ["프롬프트 / 평가", "추출 프롬프트 설계, 정답 라벨링, 정확도 측정"],
    ["데이터 / 비즈니스", "지표 분석, 경쟁 분석, 비즈니스 모델"],
]

add_table(s6, Inches(0.8), Inches(6.8), Inches(6), role_data,
          col_widths=[Inches(2), Inches(4)])


# ════════════════════════════════════════════
# S7. 마무리
# ════════════════════════════════════════════
s7 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s7, BG_DARK)

add_text(s7, Inches(0.8), Inches(2.0), Inches(11), Inches(1.5),
         "기존 도구는 회의를 기록합니다.\n우리는 회의를 완결합니다.",
         font_size=44, color=WHITE, bold=True)

add_text(s7, Inches(0.8), Inches(4.5), Inches(11), Inches(0.8),
         "Action Keeper (가칭)",
         font_size=24, color=PINK)

add_text(s7, Inches(0.8), Inches(5.3), Inches(11), Inches(0.5),
         "감사합니다.",
         font_size=20, color=LIGHT_GRAY)

# 페이지 번호 추가
for i, slide in enumerate(prs.slides):
    add_text(slide, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.3),
             str(i + 1), font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)


# ── 저장 ──
output_path = "/Users/sc.park/Desktop/0_ClaudeCode/3_KAIST/2026-Fall/1_Business_Strategy/3_TeamProject/docs/midterm_presentation.pptx"
prs.save(output_path)
print(f"PPT 생성 완료: {output_path}")
print(f"총 {len(prs.slides)}장 슬라이드")

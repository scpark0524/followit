#!/usr/bin/env python3
"""FollowIt (팔로잇) 중간발표 PPT v2 — 팀원 피드백 반영판

v1 대비 변경:
- 서비스명: Action Keeper (가칭) → FollowIt (팔로잇)
- 5단계: 연동(녹음앱 MCP) → 회의 연결 → 액션·R&R → 타임라인 → 후속관리
- 성공 기준: 구현 성공/사용 성공 2단계 + 3축 (추적 정확성·후속관리 신뢰성·실제 사용 가치)
- 수행 계획: 수직 슬라이스 — 11/7까지 최소 기능 전체 루프 → 실사용·개선 병행
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── 팔레트 (v1 실물 기준: 그린=브랜드, 코랄=페인) ──
BG = RGBColor(0x0F, 0x0F, 0x0F)          # deep charcoal
SURFACE = RGBColor(0x1A, 0x1A, 0x1A)     # card surface
SURFACE2 = RGBColor(0x22, 0x22, 0x22)    # lighter card
BORDER = RGBColor(0x2A, 0x2A, 0x2A)      # subtle border
WHITE = RGBColor(0xF5, 0xF5, 0xF5)       # off-white
GRAY1 = RGBColor(0x99, 0x99, 0x99)       # secondary text
GRAY2 = RGBColor(0x66, 0x66, 0x66)       # tertiary
BRAND = RGBColor(0x4A, 0xDE, 0x80)       # brand green
PAIN = RGBColor(0xFF, 0x6B, 0x4A)        # warm coral (문제/페인 전용)
TEAL = RGBColor(0x4E, 0xCD, 0xC4)        # secondary accent

# Slide size 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# ── 헬퍼 ──

def bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = BG

def txt(slide, l, t, w, h, text, sz=18, clr=WHITE, bold=False, align=PP_ALIGN.LEFT, name="Arial"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.color.rgb = clr
    p.font.bold = bold
    p.font.name = name
    p.alignment = align
    return box

def mtxt(slide, l, t, w, h, lines, sz=16, clr=WHITE, spacing=1.4):
    """lines: list of str or (str, {overrides})"""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        text, ov = (item, {}) if isinstance(item, str) else (item[0], item[1] if len(item) > 1 else {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(ov.get("sz", sz))
        p.font.color.rgb = ov.get("clr", clr)
        p.font.bold = ov.get("bold", False)
        p.font.name = ov.get("name", "Arial")
        p.alignment = ov.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(sz * (spacing - 1) + 2)
    return box

def card(slide, l, t, w, h, fill_clr=SURFACE, radius=Inches(0.15)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_clr
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def divider(slide, l, t, w, clr=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, Pt(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = clr
    shape.line.fill.background()
    shape.shadow.inherit = False

def pill(slide, l, t, text, clr=PAIN):
    """Small label pill"""
    w = Pt(len(text) * 14 + 30)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, Pt(28))
    shape.fill.solid()
    shape.fill.fore_color.rgb = clr
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = BG
    p.font.bold = True
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER
    return shape

def page_num(slide, n):
    txt(slide, Inches(12.5), Inches(7.05), Inches(0.6), Inches(0.3),
        str(n), sz=10, clr=GRAY2, align=PP_ALIGN.RIGHT)


# ════════════════════════════════════════════
# S1. 표지
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)

txt(s, Inches(1), Inches(1.6), Inches(6), Inches(0.4),
    "KAIST IMMS  ·  2026 Fall  ·  AI Business Strategy", sz=13, clr=GRAY2)

txt(s, Inches(1), Inches(2.3), Inches(10), Inches(1.2),
    "회의 이후를\n책임지는 AI", sz=54, clr=WHITE, bold=True)

txt(s, Inches(1), Inches(4.5), Inches(8), Inches(0.5),
    "FollowIt  ·  팔로잇", sz=24, clr=BRAND, bold=True)

divider(s, Inches(1), Inches(5.3), Inches(3))

txt(s, Inches(1), Inches(5.6), Inches(10.5), Inches(0.6),
    "회의의 맥락을 기억하고, 다음 액션을 끝까지 챙기는 AI 비서 (follow + it)", sz=16, clr=GRAY1)

txt(s, Inches(1), Inches(6.5), Inches(10), Inches(0.3),
    "TEAM_2 : 김신애 · 박승찬 · 신지엽 · 윤경혜", sz=13, clr=GRAY2)

page_num(s, 1)


# ════════════════════════════════════════════
# S2. 문제 ① — 핵심 페인
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
pill(s, Inches(1), Inches(0.7), "문제")

txt(s, Inches(1), Inches(1.3), Inches(11), Inches(1.0),
    "회의는 끝났는데, 일은 시작되지 않는다", sz=40, clr=WHITE, bold=True)

# 큰 숫자 카드
card(s, Inches(1), Inches(2.8), Inches(4.5), Inches(2.4), SURFACE)
txt(s, Inches(1.5), Inches(3.0), Inches(3.5), Inches(1.2),
    "54%", sz=72, clr=PAIN, bold=True)
txt(s, Inches(1.5), Inches(4.3), Inches(3.6), Inches(0.8),
    "회의 후 다음 할 일·담당자를 모른 채\n자리를 뜨는 직장인 비율¹", sz=14, clr=GRAY1)

# 원인 3가지
card(s, Inches(6.2), Inches(2.8), Inches(6.1), Inches(0.65), SURFACE)
txt(s, Inches(6.6), Inches(2.85), Inches(0.8), Inches(0.5),
    "①", sz=18, clr=PAIN, bold=True)
txt(s, Inches(7.2), Inches(2.85), Inches(5), Inches(0.5),
    '주인 불명확 — "그거 누가 하기로 했죠?"', sz=15, clr=WHITE)

card(s, Inches(6.2), Inches(3.65), Inches(6.1), Inches(0.65), SURFACE)
txt(s, Inches(6.6), Inches(3.7), Inches(0.8), Inches(0.5),
    "②", sz=18, clr=PAIN, bold=True)
txt(s, Inches(7.2), Inches(3.7), Inches(5), Inches(0.5),
    '기한 모호 — "다음 주쯤", "가능한 빨리"', sz=15, clr=WHITE)

card(s, Inches(6.2), Inches(4.5), Inches(6.1), Inches(0.65), SURFACE)
txt(s, Inches(6.6), Inches(4.55), Inches(0.8), Inches(0.5),
    "③", sz=18, clr=PAIN, bold=True)
txt(s, Inches(7.2), Inches(4.55), Inches(5), Inches(0.5),
    "팔로업 부재 — 다음 회의에서 확인 없이 넘어감", sz=15, clr=WHITE)

# 하단 인사이트
divider(s, Inches(1), Inches(5.8), Inches(11.3))
txt(s, Inches(1), Inches(6.05), Inches(11), Inches(0.6),
    "PM이 회의록 정리 → 태스크 등록 → 진행 확인을 수작업으로 하고 있습니다.\n이 일이 누락되는 순간, 결정사항은 사라집니다.",
    sz=15, clr=GRAY1)

txt(s, Inches(1), Inches(6.85), Inches(8), Inches(0.3),
    "1. Source: Atlassian, 지식근로자 5,000명 대상 설문 (2024)", sz=10, clr=GRAY2)

page_num(s, 2)


# ════════════════════════════════════════════
# S3. 문제 ② — 기존 서비스의 한계
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
pill(s, Inches(1), Inches(0.7), "문제")

txt(s, Inches(1), Inches(1.3), Inches(11), Inches(0.8),
    "기존 회의 AI는 '등록'에서 멈춘다", sz=40, clr=WHITE, bold=True)

# 비교 — 왼쪽: 기존, 오른쪽: 우리
card(s, Inches(1), Inches(2.5), Inches(5.5), Inches(3.4), SURFACE)
txt(s, Inches(1.4), Inches(2.65), Inches(4.5), Inches(0.4),
    "Circleback · Otter · 클로바노트 등", sz=14, clr=PAIN)
divider(s, Inches(1.4), Inches(3.2), Inches(4.7), GRAY2)

steps_existing = [
    ("녹음 → 전사", "✅"),
    ("AI 요약", "✅"),
    ("액션아이템 추출", "✅"),
    ("태스크 등록", "✅"),
    ("─── 여기서 끝 ───", ""),
]
y_pos = Inches(3.42)
for label, check in steps_existing:
    if check:
        txt(s, Inches(1.5), y_pos, Inches(3.5), Inches(0.35), label, sz=14, clr=WHITE)
        txt(s, Inches(5.2), y_pos, Inches(0.8), Inches(0.35), check, sz=14, clr=BRAND)
    else:
        txt(s, Inches(1.5), y_pos, Inches(4), Inches(0.35), label, sz=13, clr=GRAY2, align=PP_ALIGN.CENTER)
    y_pos += Inches(0.42)

# FollowIt 영역
card(s, Inches(6.8), Inches(2.5), Inches(5.5), Inches(3.4), SURFACE2)
txt(s, Inches(7.2), Inches(2.65), Inches(4.5), Inches(0.4),
    "FollowIt (팔로잇)", sz=14, clr=BRAND, bold=True)
divider(s, Inches(7.2), Inches(3.2), Inches(4.7), BRAND)

steps_ours = [
    ("기존 녹음앱 연결 (MCP/API)", "연동"),
    ("회의 연결 · 맥락 관리", "NEW"),
    ("액션 추출 + R&R 추천 + HITL 확인", "NEW"),
    ("선행관계 기반 타임라인 제안", "NEW"),
    ("팔로업 추적 + 답변 해석", "NEW"),
    ("미완료 → 다음 회의 안건", "NEW"),
]
y_pos = Inches(3.38)
for label, tag in steps_ours:
    txt(s, Inches(7.3), y_pos, Inches(3.9), Inches(0.35), label, sz=13.5, clr=WHITE)
    tag_clr = GRAY1 if tag == "연동" else BRAND
    txt(s, Inches(11.3), y_pos, Inches(0.8), Inches(0.35), tag, sz=11, clr=tag_clr, bold=(tag == "NEW"))
    y_pos += Inches(0.41)

# 핵심 메시지
txt(s, Inches(1), Inches(6.2), Inches(11.5), Inches(0.5),
    '“기존 도구는 회의를 기록합니다.  우리는 회의를 완결합니다.”',
    sz=20, clr=BRAND, bold=True)

txt(s, Inches(1), Inches(6.7), Inches(11.5), Inches(0.3),
    "녹음·전사·추출은 기존 앱 위에 올라타고, 회의 연결 · 확인 · 추적 · 완결에 집중합니다.",
    sz=13, clr=GRAY1)

page_num(s, 3)


# ════════════════════════════════════════════
# S4. 솔루션 — 5단계 워크플로
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
pill(s, Inches(1), Inches(0.7), "솔루션", TEAL)

txt(s, Inches(1), Inches(1.3), Inches(11.5), Inches(0.8),
    "회의를 기억하고, 연결하고, 끝까지 챙긴다", sz=40, clr=WHITE, bold=True)

# 5 stage flow — horizontal
stage_data = [
    ("1", "연동", "녹음앱 연결 (MCP)\n전사문 가져오기"),
    ("2", "회의 연결", "이전 회의와 묶기\n기존 액션 이어 추적"),
    ("3", "액션 · R&R", "액션 추출·R&R 추천\n모호하면 HITL 확인"),
    ("4", "타임라인", "기한·선행관계 역산\n변경사항 반영"),
    ("5", "후속관리", "진행 확인·답변 해석\n미완료 → 다음 안건"),
]

cw = Inches(2.05)
gap = Inches(0.28)
sx = Inches(0.8)

for i, (num, title, body) in enumerate(stage_data):
    x = sx + (cw + gap) * i
    c = card(s, x, Inches(2.5), cw, Inches(2.0), SURFACE)
    txt(s, x + Pt(14), Inches(2.6), Inches(0.5), Inches(0.4),
        num, sz=28, clr=BRAND, bold=True)
    txt(s, x + Pt(14), Inches(3.1), cw - Pt(28), Inches(0.4),
        title, sz=16, clr=WHITE, bold=True)
    txt(s, x + Pt(14), Inches(3.5), cw - Pt(28), Inches(0.9),
        body, sz=11.5, clr=GRAY1)
    if i < 4:
        txt(s, x + cw, Inches(3.2), gap, Inches(0.4),
            "→", sz=20, clr=GRAY2, align=PP_ALIGN.CENTER)

# 순환 표시
txt(s, Inches(2.5), Inches(4.75), Inches(8.3), Inches(0.4),
    "↻  다음 회의 녹음 → 이전 회의·액션과 자동 연결  (closed loop)",
    sz=15, clr=BRAND, bold=True, align=PP_ALIGN.CENTER)

# AI가 필요한 이유
divider(s, Inches(1), Inches(5.4), Inches(11.3))
txt(s, Inches(1), Inches(5.65), Inches(6), Inches(0.3),
    "왜 AI여야 하는가 (규칙으로는 불가능한 3가지)", sz=14, clr=WHITE, bold=True)

mtxt(s, Inches(1), Inches(6.05), Inches(11.5), Inches(1.2),
     [
         ('“그건 제가 한번 볼게요”  →  액션아이템 판단 (한국어 화법 맥락)', {"sz": 13, "clr": GRAY1}),
         ('“중간발표 전까지는 돼야죠”  →  암묵적 기한을 실제 날짜로 변환', {"sz": 13, "clr": GRAY1}),
         ('“자료가 아직 안 와서요”  →  막힘(blocker) 상태 자동 분류', {"sz": 13, "clr": GRAY1}),
     ], spacing=1.5)

page_num(s, 4)


# ════════════════════════════════════════════
# S5. 성공 기준 — 2단계 + 3축
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
pill(s, Inches(1), Inches(0.7), "성공 기준", TEAL)

txt(s, Inches(1), Inches(1.3), Inches(11.5), Inches(0.8),
    "회의 이후, 사람이 챙기는 부담이 줄었는가", sz=40, clr=WHITE, bold=True)

# 2단계 성공
card(s, Inches(1), Inches(2.4), Inches(5.55), Inches(0.9), SURFACE2)
txt(s, Inches(1.3), Inches(2.52), Inches(1.6), Inches(0.35),
    "① 구현 성공", sz=15, clr=BRAND, bold=True)
txt(s, Inches(2.9), Inches(2.52), Inches(3.5), Inches(0.7),
    "회의 가져오기 → 이전 회의 연결 → 액션 확인\n→ 후속관리가 실제로 작동하는가", sz=11.5, clr=GRAY1)

card(s, Inches(6.75), Inches(2.4), Inches(5.55), Inches(0.9), SURFACE2)
txt(s, Inches(7.05), Inches(2.52), Inches(1.6), Inches(0.35),
    "② 사용 성공", sz=15, clr=BRAND, bold=True)
txt(s, Inches(8.65), Inches(2.52), Inches(3.5), Inches(0.7),
    "쓰면서 오류·누락이 줄고, 사람이 직접\n챙기는 시간과 부담이 줄었는가", sz=11.5, clr=GRAY1)

# 3 metric cards
metrics = [
    ("추적 정확성", "연결된 회의의 액션을\n빠짐없이 잇는가",
     ["액션 누락 · 중복 건수", "담당자 · 기한 수정 건수"]),
    ("후속관리 신뢰성", "제때 확인하고,\n제대로 반영하는가",
     ["예정된 F/U 정상 처리율", "누락 · 불필요 알림 건수"]),
    ("실제 사용 가치", "도움이 되고,\n계속 쓰고 싶은가",
     ["기한 내 완료율 변화", "유용성 · 지속 사용 의향"]),
]

for i, (title, sub, lines) in enumerate(metrics):
    x = Inches(1) + Inches(3.9) * i
    card(s, x, Inches(3.6), Inches(3.6), Inches(2.1), SURFACE)
    txt(s, x + Pt(20), Inches(3.75), Inches(3), Inches(0.4),
        title, sz=19, clr=BRAND, bold=True)
    txt(s, x + Pt(20), Inches(4.15), Inches(3.1), Inches(0.55),
        sub, sz=11.5, clr=GRAY1)
    divider(s, x + Pt(20), Inches(4.75), Inches(2.8), GRAY2)
    mtxt(s, x + Pt(20), Inches(4.9), Inches(3.1), Inches(0.8),
         [(line, {"sz": 13, "clr": WHITE}) for line in lines],
         spacing=1.4)

# 실험 설계
divider(s, Inches(1), Inches(6.0), Inches(11.3))
mtxt(s, Inches(1), Inches(6.2), Inches(11.5), Inches(1),
     [("목표치는 초기 2주 실사용 결과를 보고 확정 (11/21)  ·  도입 전 수작업 관리 시간을 측정해 전후 비교", {"sz": 13, "clr": GRAY1}),
      ('교란 변수(학기말 압박) 인정 → “AI만 잡아낸 액션 수” 등 시기 무관 지표 병행', {"sz": 13, "clr": GRAY1})],
     spacing=1.3)

page_num(s, 5)


# ════════════════════════════════════════════
# S6. 수행 계획 — 수직 슬라이스
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
pill(s, Inches(1), Inches(0.7), "수행 계획", TEAL)

txt(s, Inches(1), Inches(1.3), Inches(11.5), Inches(0.8),
    "작게 전체 루프를 돌리고, 쓰면서 완성한다", sz=40, clr=WHITE, bold=True)

timeline = [
    ("10/24", "중간 발표", "문제 정의 + 설계 발표", BRAND),
    ("10/25 ~ 11/7", "최소 기능 구현", "녹음앱 1개 연결 · 회의 묶음 선택 · 액션 확인 · 기본 팔로업", WHITE),
    ("11/8 ~ 11/21", "실제 사용 시작", "팀 회의 적용 · 회의 연결 추천 · R&R·일정 보완 · 목표치 확정", WHITE),
    ("11/22 ~ 12/5", "사용하면서 개선", "답변 해석 · 막힘 파악 · 이전 업무 업데이트 · 안건 연결", WHITE),
    ("12/6 ~ 12/18", "최종 검증 · 분석", "관리 시간 · 추적 오류 · F/U 정상 처리율 · 팀원 평가 정리", WHITE),
    ("12/19", "기말 발표", "라이브 데모 + 실제 사용 결과 + 개선 사례", BRAND),
]

for i, (period, title, desc, clr) in enumerate(timeline):
    y = Inches(2.5) + Inches(0.62) * i
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y + Pt(4), Pt(10), Pt(10))
    dot.fill.solid()
    dot.fill.fore_color.rgb = clr if clr == BRAND else GRAY2
    dot.line.fill.background()
    dot.shadow.inherit = False
    if i < len(timeline) - 1:
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  Inches(1.2) + Pt(4), y + Pt(16), Pt(2), Inches(0.5))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER
        line.line.fill.background()
        line.shadow.inherit = False
    txt(s, Inches(1.8), y, Inches(2.2), Inches(0.35), period, sz=13, clr=GRAY1)
    txt(s, Inches(4.2), y, Inches(3.3), Inches(0.35), title, sz=15, clr=clr, bold=True)
    txt(s, Inches(7.6), y, Inches(5.3), Inches(0.35), desc, sz=12.5, clr=GRAY1)

# 핵심 메시지
txt(s, Inches(1), Inches(6.5), Inches(11.5), Inches(0.4),
    "11월 초부터 최소 기능을 팀 회의에 적용 → 사용·개선을 병행하며 기말까지 효과 검증",
    sz=17, clr=BRAND, bold=True)

page_num(s, 6)


# ════════════════════════════════════════════
# S7. 마무리
# ════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)

txt(s, Inches(1), Inches(2.2), Inches(11), Inches(1.8),
    "기존 도구는\n회의를 기록합니다.", sz=46, clr=GRAY1, bold=True)

txt(s, Inches(1), Inches(4.2), Inches(11), Inches(1.0),
    "우리는 회의를 완결합니다.", sz=46, clr=WHITE, bold=True)

divider(s, Inches(1), Inches(5.6), Inches(2.5), BRAND)

txt(s, Inches(1), Inches(5.9), Inches(6), Inches(0.4),
    "FollowIt  ·  팔로잇", sz=18, clr=BRAND, bold=True)

txt(s, Inches(1), Inches(6.5), Inches(5), Inches(0.3),
    "감사합니다", sz=16, clr=GRAY2)

page_num(s, 7)


# ── 저장 ──
out = "/Users/sc.park/Desktop/0_ClaudeCode/3_KAIST/2026-Fall/1_Business_Strategy/3_TeamProject/docs/midterm_presentation_v2.pptx"
prs.save(out)
print(f"✅ PPT v2 생성 완료: {out}")
print(f"   총 {len(prs.slides)}장 슬라이드")
